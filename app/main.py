from __future__ import annotations
from datetime import datetime, timedelta
import secrets
from pathlib import Path
from io import BytesIO
from urllib.parse import urlencode

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from fastapi import FastAPI, Request, Depends, BackgroundTasks, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from sqlalchemy import select, func
from sqlalchemy.orm import Session
from starlette.middleware.sessions import SessionMiddleware

from app.config import settings
from app.database import Base, engine, get_db, SessionLocal
from app.models import User, OAuthCredential, AccurateDatabase, SyncJob, GLAccount, JournalHeader, JournalLine, Department, Project
from app.core.security import hash_password, verify_password, generate_password, encrypt_secret, csrf_token, validate_csrf
from app.core.auth import current_user, is_admin, can_use_app, is_trial
from app.services.accurate_client import AccurateOAuthClient, AccurateClient, credential_access_token, client_for_database
from app.services.sync_service import run_full_sync, run_journal_sync, journal_detail_diagnostic
from app.reporting.balance_sheet import build_balance_sheet
from app.reporting.profit_loss import build_profit_loss
from app.reporting.pdf_export import balance_sheet_pdf, profit_loss_pdf

BASE_DIR = Path(__file__).resolve().parent
app = FastAPI(title=settings.app_name)
app.add_middleware(SessionMiddleware, secret_key=settings.secret_key, same_site="lax", https_only=settings.cookie_secure, max_age=60*60*24*14)
app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")
templates = Jinja2Templates(directory=BASE_DIR / "templates")


def flash(request: Request, message: str, kind: str = "info"):
    request.session.setdefault("flashes", []).append({"message": message, "kind": kind})


def template_context(request: Request, user: User | None = None, **extra):
    ctx = {
        "request": request,
        "app_name": settings.app_name,
        "user": user,
        "csrf": csrf_token(request.session),
        "flashes": request.session.pop("flashes", []),
        "is_trial": bool(user and is_trial(user)),
        "now": datetime.utcnow(),
    }
    ctx.update(extra)
    return ctx


def require_user(request: Request, db: Session) -> User:
    user = current_user(request, db)
    if not user:
        raise HTTPException(401)
    if user.role != "ADMIN" and user.status == "SUSPENDED":
        request.session.clear()
        raise HTTPException(401)
    return user


def max_databases(user: User) -> int:
    if user.role == "ADMIN":
        return settings.active_max_databases
    return settings.active_max_databases if user.status == "ACTIVE" else settings.trial_max_databases


@app.on_event("startup")
def startup():
    Base.metadata.create_all(engine)
    db = SessionLocal()
    try:
        admin = db.scalar(select(User).where(User.email == settings.admin_email))
        if not admin:
            admin = User(
                email=settings.admin_email,
                full_name="Administrator",
                password_hash=hash_password(settings.admin_password),
                role="ADMIN",
                status="ACTIVE",
                subscription_started_at=datetime.utcnow(),
            )
            db.add(admin); db.commit()
    finally:
        db.close()


@app.exception_handler(401)
async def unauthorized(request: Request, exc):
    return RedirectResponse("/login", status_code=303)


@app.get("/health")
def health():
    return {"ok": True, "app": settings.app_name}


@app.get("/", response_class=HTMLResponse)
def root(request: Request, db: Session = Depends(get_db)):
    user = current_user(request, db)
    return RedirectResponse("/dashboard" if user else "/login", status_code=303)


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request, db: Session = Depends(get_db)):
    if current_user(request, db):
        return RedirectResponse("/dashboard", status_code=303)
    return templates.TemplateResponse("login.html", template_context(request))


@app.post("/login")
async def login(request: Request, db: Session = Depends(get_db)):
    form = await request.form()
    if not validate_csrf(request.session, form.get("csrf")):
        raise HTTPException(400, "Invalid CSRF")
    email = str(form.get("email") or "").strip().lower()
    password = str(form.get("password") or "")
    user = db.scalar(select(User).where(User.email == email))
    if not user or not verify_password(password, user.password_hash):
        flash(request, "Email atau password salah.", "error")
        return RedirectResponse("/login", status_code=303)
    if user.status == "SUSPENDED":
        flash(request, "Akun Anda sedang disuspend. Hubungi administrator.", "error")
        return RedirectResponse("/login", status_code=303)
    request.session["user_id"] = user.id
    flash(request, f"Selamat datang, {user.full_name or user.email}.", "success")
    if user.must_change_password and user.role != "ADMIN":
        flash(request, "Password Anda direset admin. Silakan ganti password sebelum melanjutkan.", "warning")
        return RedirectResponse("/account", status_code=303)
    return RedirectResponse("/admin" if user.role == "ADMIN" else "/dashboard", status_code=303)


@app.get("/register", response_class=HTMLResponse)
def register_page(request: Request, db: Session = Depends(get_db)):
    if current_user(request, db):
        return RedirectResponse("/dashboard", status_code=303)
    return templates.TemplateResponse("register.html", template_context(request, trial_days=settings.trial_days))


@app.post("/register")
async def register(request: Request, db: Session = Depends(get_db)):
    form = await request.form()
    if not validate_csrf(request.session, form.get("csrf")):
        raise HTTPException(400, "Invalid CSRF")
    email = str(form.get("email") or "").strip().lower()
    full_name = str(form.get("full_name") or "").strip()
    password = str(form.get("password") or "")
    if len(password) < 8 or "@" not in email:
        flash(request, "Gunakan email valid dan password minimal 8 karakter.", "error")
        return RedirectResponse("/register", status_code=303)
    if db.scalar(select(User).where(User.email == email)):
        flash(request, "Email sudah terdaftar.", "error")
        return RedirectResponse("/register", status_code=303)
    now = datetime.utcnow()
    user = User(
        email=email, full_name=full_name, password_hash=hash_password(password),
        role="USER", status="TRIAL", trial_started_at=now,
        trial_ends_at=now + timedelta(days=settings.trial_days),
    )
    db.add(user); db.commit(); db.refresh(user)
    request.session["user_id"] = user.id
    flash(request, f"Trial {settings.trial_days} hari aktif. Selama trial hanya akun Kas/Bank yang terbuka penuh.", "success")
    return RedirectResponse("/dashboard", status_code=303)


@app.post("/logout")
async def logout(request: Request):
    form = await request.form()
    if not validate_csrf(request.session, form.get("csrf")):
        raise HTTPException(400, "Invalid CSRF")
    request.session.clear()
    return RedirectResponse("/login", status_code=303)


@app.get("/dashboard", response_class=HTMLResponse)
def dashboard(request: Request, db: Session = Depends(get_db)):
    user = require_user(request, db)
    if user.role == "ADMIN":
        return RedirectResponse("/admin", status_code=303)
    databases = db.scalars(select(AccurateDatabase).where(AccurateDatabase.user_id == user.id).order_by(AccurateDatabase.selected.desc(), AccurateDatabase.alias)).all()
    selected = next((d for d in databases if d.selected), databases[0] if databases else None)
    last_job = db.scalar(select(SyncJob).where(SyncJob.user_id == user.id).order_by(SyncJob.id.desc()).limit(1))
    days_left = None
    if user.status == "TRIAL" and user.trial_ends_at:
        days_left = max(0, (user.trial_ends_at.date() - datetime.utcnow().date()).days)
    return templates.TemplateResponse("dashboard.html", template_context(
        request, user, databases=databases, selected=selected, last_job=last_job,
        max_databases=max_databases(user), days_left=days_left,
    ))


@app.get("/account", response_class=HTMLResponse)
def account_page(request: Request, db: Session = Depends(get_db)):
    user = require_user(request, db)
    return templates.TemplateResponse("account.html", template_context(request, user))


@app.post("/account/password")
async def change_password(request: Request, db: Session = Depends(get_db)):
    user = require_user(request, db)
    form = await request.form()
    if not validate_csrf(request.session, form.get("csrf")):
        raise HTTPException(400, "Invalid CSRF")
    current = str(form.get("current_password") or "")
    new = str(form.get("new_password") or "")
    if not verify_password(current, user.password_hash) or len(new) < 8:
        flash(request, "Password lama salah atau password baru terlalu pendek.", "error")
    else:
        user.password_hash = hash_password(new); user.must_change_password = False; db.commit()
        flash(request, "Password berhasil diubah.", "success")
    return RedirectResponse("/account", status_code=303)


# ---------- Accurate Online OAuth ----------
@app.get("/accurate/connect")
def accurate_connect(request: Request, db: Session = Depends(get_db)):
    user = require_user(request, db)
    if not can_use_app(user):
        flash(request, "Aktifkan langganan untuk menghubungkan Accurate Online.", "error")
        return RedirectResponse("/dashboard", status_code=303)
    if not settings.aol_client_id or not settings.aol_client_secret:
        flash(request, "AOL_CLIENT_ID / AOL_CLIENT_SECRET belum dikonfigurasi di server.", "error")
        return RedirectResponse("/dashboard", status_code=303)
    state = secrets.token_urlsafe(24)
    request.session["oauth_state"] = state
    return RedirectResponse(AccurateOAuthClient().authorization_url(state), status_code=302)


@app.get("/accurate/oauth/callback")
def accurate_callback(request: Request, code: str = "", state: str = "", error: str = "", db: Session = Depends(get_db)):
    user = require_user(request, db)
    if error:
        flash(request, f"OAuth Accurate gagal: {error}", "error"); return RedirectResponse("/dashboard", status_code=303)
    expected = request.session.pop("oauth_state", "")
    if not expected or not secrets.compare_digest(expected, state):
        flash(request, "OAuth state tidak valid.", "error"); return RedirectResponse("/dashboard", status_code=303)
    try:
        data = AccurateOAuthClient().exchange_code(code)
        expires = datetime.utcnow() + timedelta(seconds=int(data.get("expires_in") or 1295999))
        cred = user.oauth or OAuthCredential(user_id=user.id, access_token_enc="")
        if not user.oauth: db.add(cred)
        cred.access_token_enc = encrypt_secret(data.get("access_token", ""))
        cred.refresh_token_enc = encrypt_secret(data.get("refresh_token", ""))
        cred.token_type = data.get("token_type", "bearer"); cred.scope = data.get("scope", ""); cred.expires_at = expires
        aol_user = data.get("user") or {}
        cred.aol_user_email = str(aol_user.get("email") or aol_user.get("nickname") or "")
        cred.aol_user_name = str(aol_user.get("name") or "")
        db.commit()
        flash(request, "Accurate Online berhasil terhubung. Pilih database yang ingin digunakan.", "success")
        return RedirectResponse("/accurate/databases", status_code=303)
    except Exception as exc:
        flash(request, f"Gagal menukar OAuth code: {exc}", "error")
        return RedirectResponse("/dashboard", status_code=303)


@app.get("/accurate/databases", response_class=HTMLResponse)
def accurate_databases(request: Request, db: Session = Depends(get_db)):
    user = require_user(request, db)
    if not user.oauth:
        return RedirectResponse("/accurate/connect", status_code=303)
    try:
        token = credential_access_token(user.oauth, db)
        available = AccurateClient(token).db_list()
    except Exception as exc:
        flash(request, f"Gagal membaca daftar database AOL: {exc}", "error")
        available = []
    saved = db.scalars(select(AccurateDatabase).where(AccurateDatabase.user_id == user.id)).all()
    saved_ids = {d.accurate_db_id for d in saved}
    return templates.TemplateResponse("accurate_databases.html", template_context(
        request, user, available=available, saved=saved, saved_ids=saved_ids,
        max_databases=max_databases(user),
    ))


@app.post("/accurate/databases/add")
async def add_database(request: Request, db: Session = Depends(get_db)):
    user = require_user(request, db)
    form = await request.form()
    if not validate_csrf(request.session, form.get("csrf")): raise HTTPException(400, "Invalid CSRF")
    count = db.scalar(select(func.count()).select_from(AccurateDatabase).where(AccurateDatabase.user_id == user.id)) or 0
    if count >= max_databases(user):
        flash(request, f"Batas database akun Anda adalah {max_databases(user)} database.", "error")
        return RedirectResponse("/accurate/databases", status_code=303)
    aid = int(form.get("accurate_db_id"))
    alias = str(form.get("alias") or f"Database {aid}")
    if db.scalar(select(AccurateDatabase).where(AccurateDatabase.user_id == user.id, AccurateDatabase.accurate_db_id == aid)):
        flash(request, "Database tersebut sudah ditambahkan.", "info")
        return RedirectResponse("/accurate/databases", status_code=303)
    try:
        token = credential_access_token(user.oauth, db)
        opened = AccurateClient(token).open_db(aid)
        if not opened.get("session") or not opened.get("host"):
            raise RuntimeError(str(opened))
        db.execute(AccurateDatabase.__table__.update().where(AccurateDatabase.user_id == user.id).values(selected=False))
        row = AccurateDatabase(user_id=user.id, accurate_db_id=aid, alias=alias, host=opened["host"], session_enc=encrypt_secret(opened["session"]), selected=True)
        db.add(row); db.commit()
        flash(request, f"{alias} berhasil ditambahkan dan dipilih.", "success")
    except Exception as exc:
        db.rollback(); flash(request, f"Gagal membuka database AOL: {exc}", "error")
    return RedirectResponse("/accurate/databases", status_code=303)


@app.post("/accurate/databases/{database_id}/select")
async def select_database(database_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_user(request, db); form = await request.form()
    if not validate_csrf(request.session, form.get("csrf")): raise HTTPException(400, "Invalid CSRF")
    row = db.scalar(select(AccurateDatabase).where(AccurateDatabase.id == database_id, AccurateDatabase.user_id == user.id))
    if row:
        db.execute(AccurateDatabase.__table__.update().where(AccurateDatabase.user_id == user.id).values(selected=False)); row.selected = True; db.commit()
    return RedirectResponse("/dashboard", status_code=303)


@app.post("/accurate/databases/{database_id}/remove")
async def remove_database(database_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_user(request, db); form = await request.form()
    if not validate_csrf(request.session, form.get("csrf")): raise HTTPException(400, "Invalid CSRF")
    row = db.scalar(select(AccurateDatabase).where(AccurateDatabase.id == database_id, AccurateDatabase.user_id == user.id))
    if row:
        db.delete(row); db.commit(); flash(request, "Database dihapus dari SUPERTOOLS. Data di Accurate tidak berubah.", "success")
    return RedirectResponse("/accurate/databases", status_code=303)


@app.post("/sync/{database_id}")
async def start_sync(database_id: int, request: Request, background_tasks: BackgroundTasks, db: Session = Depends(get_db)):
    user = require_user(request, db); form = await request.form()
    if not validate_csrf(request.session, form.get("csrf")): raise HTTPException(400, "Invalid CSRF")
    row = db.scalar(select(AccurateDatabase).where(AccurateDatabase.id == database_id, AccurateDatabase.user_id == user.id))
    if not row or not user.oauth:
        flash(request, "Database/OAuth belum siap.", "error"); return RedirectResponse("/dashboard", status_code=303)
    running = db.scalar(select(SyncJob).where(SyncJob.database_id == row.id, SyncJob.status.in_(["QUEUED","RUNNING"])).order_by(SyncJob.id.desc()).limit(1))
    if running:
        flash(request, "Sync masih berjalan.", "info"); return RedirectResponse("/dashboard", status_code=303)
    job = SyncJob(user_id=user.id, database_id=row.id, status="QUEUED", message="Menunggu worker...")
    db.add(job); db.commit(); db.refresh(job)
    background_tasks.add_task(run_full_sync, job.id)
    flash(request, "Sync dimulai. Halaman dapat direfresh untuk melihat progress.", "success")
    return RedirectResponse("/dashboard", status_code=303)


@app.get("/api/sync/{job_id}")
def sync_status(job_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_user(request, db)
    job = db.scalar(select(SyncJob).where(SyncJob.id == job_id, SyncJob.user_id == user.id))
    if not job: raise HTTPException(404)
    return {"status":job.status,"progress":job.progress,"message":job.message}



@app.get("/journals", response_class=HTMLResponse)
def journals_page(request: Request, db_id: int | None = None, account: str = "", department: str = "", project: str = "", diagnose: int = 0, db: Session = Depends(get_db)):
    user = require_user(request, db)
    q = select(AccurateDatabase).where(AccurateDatabase.user_id == user.id)
    q = q.where(AccurateDatabase.id == db_id) if db_id else q.where(AccurateDatabase.selected == True)  # noqa: E712
    database = db.scalar(q)
    if not database:
        flash(request, "Pilih database Accurate Online terlebih dahulu.", "info")
        return RedirectResponse("/accurate/databases", status_code=303)

    header_count = db.scalar(select(func.count()).select_from(JournalHeader).where(JournalHeader.database_id == database.id)) or 0
    line_count = db.scalar(select(func.count()).select_from(JournalLine).where(JournalLine.database_id == database.id)) or 0
    dept_count = db.scalar(select(func.count()).select_from(JournalLine).where(JournalLine.database_id == database.id, JournalLine.department_name != "")) or 0
    project_count = db.scalar(select(func.count()).select_from(JournalLine).where(JournalLine.database_id == database.id, JournalLine.project_no != "")) or 0
    cash_count = db.scalar(
        select(func.count()).select_from(JournalLine)
        .join(GLAccount, (GLAccount.database_id == JournalLine.database_id) & (GLAccount.account_no == JournalLine.account_no))
        .where(JournalLine.database_id == database.id, GLAccount.account_type == "CASH_BANK")
    ) or 0
    unknown_accounts = db.scalar(
        select(func.count()).select_from(JournalLine)
        .outerjoin(GLAccount, (GLAccount.database_id == JournalLine.database_id) & (GLAccount.account_no == JournalLine.account_no))
        .where(JournalLine.database_id == database.id, GLAccount.id.is_(None))
    ) or 0

    stmt = (
        select(JournalLine, JournalHeader, GLAccount)
        .join(JournalHeader, (JournalHeader.database_id == JournalLine.database_id) & (JournalHeader.accurate_id == JournalLine.journal_accurate_id))
        .outerjoin(GLAccount, (GLAccount.database_id == JournalLine.database_id) & (GLAccount.account_no == JournalLine.account_no))
        .where(JournalLine.database_id == database.id)
        .order_by(JournalHeader.trans_date.desc(), JournalHeader.accurate_id.desc(), JournalLine.line_accurate_id.asc())
    )
    if account:
        stmt = stmt.where(JournalLine.account_no.ilike(f"%{account}%"))
    if department:
        stmt = stmt.where(JournalLine.department_name.ilike(f"%{department}%"))
    if project:
        stmt = stmt.where(JournalLine.project_no.ilike(f"%{project}%"))
    rows = db.execute(stmt.limit(1000)).all()
    last_job = db.scalar(select(SyncJob).where(SyncJob.database_id == database.id).order_by(SyncJob.id.desc()).limit(1))

    api_diagnostic = None
    diagnostic_error = ""
    if diagnose and user.oauth:
        try:
            client = client_for_database(database, user.oauth, db)
            first_headers = list(client.paged_list("journal-voucher", fields="id,number,transDate,description", page_size=1))
            if first_headers:
                api_diagnostic = journal_detail_diagnostic(client, first_headers[0])
            else:
                diagnostic_error = "journal-voucher/list.do tidak mengembalikan header."
        except Exception as exc:
            diagnostic_error = f"{type(exc).__name__}: {exc}"

    return templates.TemplateResponse("journals.html", template_context(
        request, user, database=database, rows=rows, header_count=header_count, line_count=line_count,
        dept_count=dept_count, project_count=project_count, cash_count=cash_count, unknown_accounts=unknown_accounts,
        last_job=last_job, account_filter=account, department_filter=department, project_filter=project,
        api_diagnostic=api_diagnostic, diagnostic_error=diagnostic_error,
    ))


@app.post("/journals/load-all")
async def journals_load_all(request: Request, background_tasks: BackgroundTasks, db: Session = Depends(get_db)):
    user = require_user(request, db)
    form = await request.form()
    if not validate_csrf(request.session, form.get("csrf")):
        raise HTTPException(400, "Invalid CSRF")
    database_id = int(form.get("database_id") or 0)
    row = db.scalar(select(AccurateDatabase).where(AccurateDatabase.id == database_id, AccurateDatabase.user_id == user.id))
    if not row or not user.oauth:
        flash(request, "Database/OAuth belum siap.", "error")
        return RedirectResponse("/journals", status_code=303)
    running = db.scalar(select(SyncJob).where(SyncJob.database_id == row.id, SyncJob.status.in_(["QUEUED","RUNNING"])).order_by(SyncJob.id.desc()).limit(1))
    if running:
        flash(request, "Masih ada proses sync yang berjalan.", "info")
        return RedirectResponse(f"/journals?db_id={row.id}", status_code=303)
    job = SyncJob(user_id=user.id, database_id=row.id, status="QUEUED", message="Load All Jurnal menunggu worker...")
    db.add(job); db.commit(); db.refresh(job)
    background_tasks.add_task(run_journal_sync, job.id)
    flash(request, "Load All Jurnal dimulai. Refresh halaman ini untuk melihat jumlah jurnal/detail yang berhasil masuk.", "success")
    return RedirectResponse(f"/journals?db_id={row.id}", status_code=303)

def _selected_database(db: Session, user: User, db_id: int | None):
    q = select(AccurateDatabase).where(AccurateDatabase.user_id == user.id)
    q = q.where(AccurateDatabase.id == db_id) if db_id else q.where(AccurateDatabase.selected == True)  # noqa: E712
    return db.scalar(q)


def _dimension_filter_context(db: Session, database: AccurateDatabase, request: Request):
    departments = db.scalars(select(Department).where(Department.database_id == database.id, Department.suspended == False).order_by(Department.name)).all()  # noqa: E712
    projects = db.scalars(select(Project).where(Project.database_id == database.id, Project.suspended == False).order_by(Project.project_no, Project.name)).all()  # noqa: E712
    selected_departments = [x for x in request.query_params.getlist("department") if x]
    selected_projects = [x for x in request.query_params.getlist("project") if x]
    has_unmapped_department = bool(db.scalar(select(func.count()).select_from(JournalLine).where(JournalLine.database_id == database.id, (JournalLine.department_name == "") | JournalLine.department_name.is_(None))))
    has_unmapped_project = bool(db.scalar(select(func.count()).select_from(JournalLine).where(JournalLine.database_id == database.id, (JournalLine.project_no == "") | JournalLine.project_no.is_(None))))
    return {
        "departments": departments, "projects": projects,
        "selected_departments": selected_departments, "selected_projects": selected_projects,
        "has_unmapped_department": has_unmapped_department, "has_unmapped_project": has_unmapped_project,
    }


def _filters_query(selected_departments: list[str], selected_projects: list[str]) -> str:
    pairs=[]
    pairs.extend(("department", x) for x in selected_departments)
    pairs.extend(("project", x) for x in selected_projects)
    return urlencode(pairs)


# ---------- Reports ----------
@app.get("/report", response_class=HTMLResponse)
def report(request: Request, dimension: str = "department", as_of: str = "", db_id: int | None = None, db: Session = Depends(get_db)):
    user = require_user(request, db)
    if not can_use_app(user):
        flash(request, "Trial/langganan sudah berakhir atau akun disuspend.", "error")
        return RedirectResponse("/dashboard", status_code=303)
    database = _selected_database(db, user, db_id)
    if not database:
        flash(request, "Pilih database Accurate Online terlebih dahulu.", "info")
        return RedirectResponse("/accurate/databases", status_code=303)
    fctx = _dimension_filter_context(db, database, request)
    if dimension == "department": fctx["selected_projects"] = []
    elif dimension == "project": fctx["selected_departments"] = []
    report_data = build_balance_sheet(
        db, user, database, dimension=dimension, as_of=as_of or None,
        department_filters=fctx["selected_departments"], project_filters=fctx["selected_projects"],
    )
    fctx["filter_query"] = _filters_query(fctx["selected_departments"], fctx["selected_projects"])
    return templates.TemplateResponse("report.html", template_context(
        request, user, database=database, report=report_data, dimension=dimension, as_of=as_of, **fctx
    ))


@app.get("/export.xlsx")
def export_xlsx(request: Request, dimension: str = "department", as_of: str = "", db_id: int | None = None, db: Session = Depends(get_db)):
    user = require_user(request, db)
    database = _selected_database(db, user, db_id)
    if not database:
        raise HTTPException(404, "Database not selected")
    fctx = _dimension_filter_context(db, database, request)
    if dimension == "department": fctx["selected_projects"] = []
    elif dimension == "project": fctx["selected_departments"] = []
    data = build_balance_sheet(
        db, user, database, dimension=dimension, as_of=as_of or None,
        department_filters=fctx["selected_departments"], project_filters=fctx["selected_projects"],
    )
    wb = Workbook(); ws = wb.active; ws.title = "Balance Sheet"
    ws.append([settings.app_name]); ws.append([database.alias, "As of", as_of or "All data", "Mode", dimension])
    headers = ["Account No", "Account Name", "Account Type"] + data["dimensions"] + ["TOTAL"]
    ws.append(headers)
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="075C59"); c.alignment = Alignment(horizontal="center")
    for section in data["sections"]:
        ws.append([section["major"]]); ws.cell(ws.max_row,1).font=Font(bold=True,color="FFFFFF"); ws.cell(ws.max_row,1).fill=PatternFill("solid",fgColor="075C59")
        ws.append([section["group"]]); ws.cell(ws.max_row,1).font=Font(bold=True,color="075C59"); ws.cell(ws.max_row,1).fill=PatternFill("solid",fgColor="DFF1EF")
        for row in section["rows"]:
            vals=[]; total=0.0
            for d in data["dimensions"]:
                v=row["values"][d]; vals.append("TRIAL LOCKED" if v is None else float(v))
                if v is not None: total += float(v)
            ws.append([row["account_no"], row["name"], row["account_type"], *vals, "TRIAL LOCKED" if row.get("masked") else total])
        subtotal=[]; subtotal_total=0.0; locked=False
        for d in data["dimensions"]:
            v=section["subtotal"][d]
            if v is None: subtotal.append("TRIAL LOCKED"); locked=True
            else: subtotal.append(float(v)); subtotal_total+=float(v)
        ws.append(["", section["subtotal_label"], "", *subtotal, "TRIAL LOCKED" if locked else subtotal_total])
        for c in ws[ws.max_row]: c.font=Font(bold=True,color="075C59"); c.fill=PatternFill("solid",fgColor="EFF8F6")
    ws.append([])
    for label,key in [("TOTAL ASET","ASSET"),("TOTAL LIABILITAS","LIABILITY"),("TOTAL EKUITAS","EQUITY"),("TOTAL LIABILITAS DAN EKUITAS","PASIVA"),("BALANCE CHECK","CHECK")]:
        vals=[]; total=0.0; locked=False
        for d in data["dimensions"]:
            v=data["totals"][d][key]
            if v is None: vals.append("TRIAL LOCKED"); locked=True
            else: vals.append(float(v)); total+=float(v)
        ws.append(["",label,"",*vals,"TRIAL LOCKED" if locked else total])
        for c in ws[ws.max_row]: c.font=Font(bold=True)
    ws.freeze_panes="D4"
    for col in ws.columns:
        letter=col[0].column_letter; ws.column_dimensions[letter].width=min(max(12,max(len(str(c.value or "")) for c in col)+2),34)
    bio=BytesIO(); wb.save(bio); bio.seek(0)
    filename=f"SUPERTOOLS_{database.alias.replace(' ','_')}_{dimension}.xlsx"
    return StreamingResponse(bio, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition":f'attachment; filename="{filename}"'})


@app.get("/report.pdf")
def report_pdf(request: Request, dimension: str = "department", as_of: str = "", db_id: int | None = None, db: Session = Depends(get_db)):
    user = require_user(request, db)
    database = _selected_database(db, user, db_id)
    if not database:
        raise HTTPException(404, "Database not selected")
    fctx = _dimension_filter_context(db, database, request)
    if dimension == "department": fctx["selected_projects"] = []
    elif dimension == "project": fctx["selected_departments"] = []
    data = build_balance_sheet(
        db, user, database, dimension=dimension, as_of=as_of or None,
        department_filters=fctx["selected_departments"], project_filters=fctx["selected_projects"],
    )
    bio = balance_sheet_pdf(
        settings.app_name, database.alias, data, dimension, as_of,
        fctx["selected_departments"], fctx["selected_projects"],
    )
    filename=f"SUPERTOOLS_BS_{database.alias.replace(' ','_')}_{dimension}.pdf"
    return StreamingResponse(bio, media_type="application/pdf", headers={"Content-Disposition":f'attachment; filename="{filename}"'})


@app.get("/profit-loss", response_class=HTMLResponse)
def profit_loss(request: Request, dimension: str = "department", date_from: str = "", date_to: str = "", db_id: int | None = None, db: Session = Depends(get_db)):
    user = require_user(request, db)
    if not can_use_app(user):
        flash(request, "Trial/langganan sudah berakhir atau akun disuspend.", "error")
        return RedirectResponse("/dashboard", status_code=303)
    database = _selected_database(db, user, db_id)
    if not database:
        flash(request, "Pilih database Accurate Online terlebih dahulu.", "info")
        return RedirectResponse("/accurate/databases", status_code=303)
    fctx = _dimension_filter_context(db, database, request)
    if dimension == "department": fctx["selected_projects"] = []
    elif dimension == "project": fctx["selected_departments"] = []
    data = build_profit_loss(
        db, user, database, dimension=dimension, date_from=date_from or None, date_to=date_to or None,
        department_filters=fctx["selected_departments"], project_filters=fctx["selected_projects"],
    )
    fctx["filter_query"] = _filters_query(fctx["selected_departments"], fctx["selected_projects"])
    return templates.TemplateResponse("profit_loss.html", template_context(
        request, user, database=database, report=data, dimension=dimension, date_from=date_from, date_to=date_to, **fctx
    ))


@app.get("/profit-loss.xlsx")
def profit_loss_xlsx(request: Request, dimension: str = "department", date_from: str = "", date_to: str = "", db_id: int | None = None, db: Session = Depends(get_db)):
    user = require_user(request, db)
    database = _selected_database(db, user, db_id)
    if not database: raise HTTPException(404, "Database not selected")
    fctx = _dimension_filter_context(db, database, request)
    if dimension == "department": fctx["selected_projects"] = []
    elif dimension == "project": fctx["selected_departments"] = []
    data = build_profit_loss(db,user,database,dimension=dimension,date_from=date_from or None,date_to=date_to or None,department_filters=fctx["selected_departments"],project_filters=fctx["selected_projects"])
    wb=Workbook(); ws=wb.active; ws.title="Profit Loss"
    ws.append([settings.app_name]); ws.append([database.alias,"From",date_from or "All","To",date_to or "All","Mode",dimension])
    headers=["Account No","Account Name","Account Type"]+data["dimensions"]+["TOTAL"]; ws.append(headers)
    for c in ws[3]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="075C59")
    for section in data["sections"]:
        ws.append([section["label"]]); ws.cell(ws.max_row,1).font=Font(bold=True,color="FFFFFF"); ws.cell(ws.max_row,1).fill=PatternFill("solid",fgColor="075C59")
        for row in section["rows"]:
            vals=[]; total=0.0
            for d in data["dimensions"]:
                v=row["values"][d]; vals.append("TRIAL LOCKED" if v is None else float(v))
                if v is not None: total+=float(v)
            ws.append([row["account_no"],row["name"],row["account_type"],*vals,"TRIAL LOCKED" if row.get("masked") else total])
        vals=[]; total=0.0; locked=False
        for d in data["dimensions"]:
            v=section["subtotal"][d]
            if v is None: vals.append("TRIAL LOCKED"); locked=True
            else: vals.append(float(v)); total+=float(v)
        ws.append(["",f"TOTAL {section['label']}","",*vals,"TRIAL LOCKED" if locked else total])
        for c in ws[ws.max_row]: c.font=Font(bold=True); c.fill=PatternFill("solid",fgColor="EFF8F6")
    ws.append([])
    for label,key in [("LABA KOTOR","GROSS_PROFIT"),("LABA USAHA","OPERATING_PROFIT"),("LABA BERSIH","NET_PROFIT")]:
        vals=[]; total=0.0; locked=False
        for d in data["dimensions"]:
            v=data["summaries"][d][key]
            if v is None: vals.append("TRIAL LOCKED"); locked=True
            else: vals.append(float(v)); total+=float(v)
        ws.append(["",label,"",*vals,"TRIAL LOCKED" if locked else total])
        for c in ws[ws.max_row]: c.font=Font(bold=True)
    for col in ws.columns:
        letter=col[0].column_letter; ws.column_dimensions[letter].width=min(max(12,max(len(str(c.value or "")) for c in col)+2),34)
    bio=BytesIO(); wb.save(bio); bio.seek(0)
    filename=f"SUPERTOOLS_PL_{database.alias.replace(' ','_')}_{dimension}.xlsx"
    return StreamingResponse(bio,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="{filename}"'})


@app.get("/profit-loss.pdf")
def profit_loss_pdf_route(request: Request, dimension: str = "department", date_from: str = "", date_to: str = "", db_id: int | None = None, db: Session = Depends(get_db)):
    user = require_user(request, db)
    database = _selected_database(db, user, db_id)
    if not database:
        raise HTTPException(404, "Database not selected")
    fctx = _dimension_filter_context(db, database, request)
    if dimension == "department": fctx["selected_projects"] = []
    elif dimension == "project": fctx["selected_departments"] = []
    data = build_profit_loss(
        db, user, database, dimension=dimension, date_from=date_from or None, date_to=date_to or None,
        department_filters=fctx["selected_departments"], project_filters=fctx["selected_projects"],
    )
    bio = profit_loss_pdf(
        settings.app_name, database.alias, data, dimension, date_from, date_to,
        fctx["selected_departments"], fctx["selected_projects"],
    )
    filename=f"SUPERTOOLS_PL_{database.alias.replace(' ','_')}_{dimension}.pdf"
    return StreamingResponse(bio, media_type="application/pdf", headers={"Content-Disposition":f'attachment; filename="{filename}"'})


# ---------- Admin ----------
@app.get("/admin", response_class=HTMLResponse)
def admin_home(request: Request, db: Session = Depends(get_db)):
    admin = require_user(request, db)
    if not is_admin(admin): raise HTTPException(403)
    users = db.scalars(select(User).order_by(User.created_at.desc())).all()
    counts = {status: db.scalar(select(func.count()).select_from(User).where(User.status == status, User.role == "USER")) or 0 for status in ("TRIAL","ACTIVE","SUSPENDED","EXPIRED")}
    return templates.TemplateResponse("admin/index.html", template_context(request, admin, users=users, counts=counts))


@app.get("/admin/users/{user_id}", response_class=HTMLResponse)
def admin_user(user_id: int, request: Request, db: Session = Depends(get_db)):
    admin = require_user(request, db)
    if not is_admin(admin): raise HTTPException(403)
    target = db.get(User, user_id)
    if not target: raise HTTPException(404)
    databases = db.scalars(select(AccurateDatabase).where(AccurateDatabase.user_id == target.id)).all()
    return templates.TemplateResponse("admin/user.html", template_context(request, admin, target=target, databases=databases))


@app.post("/admin/users/{user_id}/activate")
async def admin_activate(user_id: int, request: Request, db: Session = Depends(get_db)):
    admin = require_user(request, db)
    if not is_admin(admin): raise HTTPException(403)
    form = await request.form()
    if not validate_csrf(request.session, form.get("csrf")): raise HTTPException(400, "Invalid CSRF")
    target = db.get(User, user_id); days = int(form.get("days") or 30)
    if target:
        now=datetime.utcnow(); target.status="ACTIVE"; target.subscription_started_at=now; target.subscription_ends_at=now+timedelta(days=max(1,days)); db.commit()
        flash(request, f"{target.email} diaktifkan selama {days} hari.", "success")
    return RedirectResponse(f"/admin/users/{user_id}", status_code=303)


@app.post("/admin/users/{user_id}/suspend")
async def admin_suspend(user_id: int, request: Request, db: Session = Depends(get_db)):
    admin=require_user(request,db)
    if not is_admin(admin): raise HTTPException(403)
    form=await request.form()
    if not validate_csrf(request.session, form.get("csrf")): raise HTTPException(400,"Invalid CSRF")
    target=db.get(User,user_id)
    if target and target.role!="ADMIN": target.status="SUSPENDED"; db.commit(); flash(request,f"{target.email} disuspend.","success")
    return RedirectResponse(f"/admin/users/{user_id}",status_code=303)


@app.post("/admin/users/{user_id}/trial-extend")
async def admin_trial_extend(user_id:int,request:Request,db:Session=Depends(get_db)):
    admin=require_user(request,db)
    if not is_admin(admin): raise HTTPException(403)
    form=await request.form()
    if not validate_csrf(request.session,form.get("csrf")): raise HTTPException(400,"Invalid CSRF")
    target=db.get(User,user_id); days=int(form.get("days") or settings.trial_days)
    if target:
        base=max(target.trial_ends_at or datetime.utcnow(),datetime.utcnow()); target.trial_ends_at=base+timedelta(days=max(1,days)); target.status="TRIAL"; db.commit(); flash(request,f"Trial ditambah {days} hari.","success")
    return RedirectResponse(f"/admin/users/{user_id}",status_code=303)


@app.post("/admin/users/{user_id}/reset-password")
async def admin_reset_password(user_id:int,request:Request,db:Session=Depends(get_db)):
    admin=require_user(request,db)
    if not is_admin(admin): raise HTTPException(403)
    form=await request.form()
    if not validate_csrf(request.session,form.get("csrf")): raise HTTPException(400,"Invalid CSRF")
    target=db.get(User,user_id)
    if target:
        temp=generate_password(); target.password_hash=hash_password(temp); target.must_change_password=True; db.commit(); flash(request,f"Temporary password untuk {target.email}: {temp}","warning")
    return RedirectResponse(f"/admin/users/{user_id}",status_code=303)
